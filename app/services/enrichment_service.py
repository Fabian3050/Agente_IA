import asyncio
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, and_, case
from datetime import datetime
from typing import Optional, List, Dict, Any
from sqlalchemy.orm import Session

from app.models.db_models import DSpaceEnrichmentQueue
from app.models.metadata_model import MetadataResponse
from app.services.metadata_service import MetadataService
from app.services.classification_service import ClassificationService

class EnrichmentService:
    """
    Servicio encargado de enriquecer las publicaciones de la cola DSpaceEnrichmentQueue.
    Fase 1: Consulta a fuentes académicas (OpenAlex, CrossRef, Semantic Scholar, Europe PMC, etc.)
            para recuperar abstract, palabras clave, fuentes de financiamiento y derechos de acceso.
    Fase 2: Clasificación temática con IA (Ollama / LLaMA 3.2) para asignar áreas OCDE y ODS.
    """

    def __init__(self):
        self.metadata_service = MetadataService()
        self.classification_service = ClassificationService()

    async def enrich_item(self, item: DSpaceEnrichmentQueue, db: Session, classify_with_ai: bool = True) -> Dict[str, Any]:
        """
        Ejecuta el pipeline completo de enriquecimiento para un registro individual.
        """
        item.estado = "EN_PROCESO"
        item.intentos += 1
        item.updated_at = datetime.utcnow()
        db.commit()

        fuentes_utilizadas = set()
        abstract_encontrado: Optional[str] = None
        keywords_acumuladas = set()
        financiamiento_acumulado = set()
        derechos_acceso: Optional[str] = None
        doi_encontrado: Optional[str] = item.doi
        raw_responses: List[Dict[str, Any]] = []

        try:
            # ----------------------------------------------------
            # FASE 1: Búsqueda en Fuentes Académicas
            # ----------------------------------------------------
            resultados: List[MetadataResponse] = []

            # 1.1 Si tiene DOI, consultamos directamente por DOI
            if item.doi and item.doi.strip():
                try:
                    resultados = await self.metadata_service.get_metadata_by_doi(item.doi.strip())
                except Exception as e:
                    print(f"Error consultando por DOI {item.doi}: {e}")

            # 1.2 Si no tiene DOI o la consulta por DOI no devolvió resultados, buscamos por título
            if not resultados and item.titulo and len(item.titulo.strip()) > 5:
                try:
                    # Búsqueda primero en OpenAlex (optimizado para búsqueda difusa)
                    candidatos = await self.metadata_service.search_metadata(item.titulo.strip(), source="openalex", limit_per_source=3)
                    if candidatos:
                        resultados.extend(candidatos)
                    else:
                        # Si no hay en OpenAlex, probamos en todas las fuentes
                        resultados = await self.metadata_service.search_metadata(item.titulo.strip(), limit_per_source=2)
                except Exception as e:
                    print(f"Error consultando por título '{item.titulo}': {e}")

            # 1.3 Procesar y combinar los metadatos obtenidos de las fuentes
            if resultados:
                for r in resultados:
                    if r.source:
                        fuentes_utilizadas.add(r.source)
                    
                    # Abstract (tomamos el primero no vacío con longitud significativa)
                    if not abstract_encontrado and r.abstract and len(r.abstract.strip()) > 30:
                        abstract_encontrado = r.abstract.strip()

                    # DOI (si el ítem no tenía DOI y la fuente nos lo proporciona)
                    if not doi_encontrado and r.doi:
                        doi_encontrado = r.doi.strip()

                    # Palabras clave
                    if r.keywords:
                        for kw in r.keywords:
                            if kw and kw.strip():
                                keywords_acumuladas.add(kw.strip())

                    # Financiamiento
                    if r.funding_source:
                        for fn in r.funding_source:
                            if fn and fn.strip():
                                financiamiento_acumulado.add(fn.strip())

                    # Derechos de acceso (Open Access, Closed, etc.)
                    if not derechos_acceso and r.derechos_acceso:
                        derechos_acceso = r.derechos_acceso.strip()

                    # Guardar respuesta resumida para respaldo
                    raw_responses.append({
                        "source": r.source,
                        "title": r.title,
                        "doi": r.doi,
                        "has_abstract": bool(r.abstract),
                        "keywords_count": len(r.keywords) if r.keywords else 0,
                        "derechos_acceso": r.derechos_acceso
                    })

                # Actualizar campos del registro
                if abstract_encontrado:
                    item.abstract_enriquecido = abstract_encontrado
                if keywords_acumuladas:
                    item.palabras_clave = list(keywords_acumuladas)
                if financiamiento_acumulado:
                    item.fuente_financiamiento = list(financiamiento_acumulado)
                if derechos_acceso:
                    item.derechos_acceso = derechos_acceso
                if doi_encontrado and not item.doi:
                    item.doi = doi_encontrado

                item.fuente_enriquecimiento = ", ".join(sorted(fuentes_utilizadas)) if fuentes_utilizadas else None
                item.metadatos_recuperados = {"fuentes_consultadas": list(fuentes_utilizadas), "respuestas": raw_responses}
                item.estado = "ENRIQUECIDO_FUENTES"
            else:
                item.estado = "NO_ENCONTRADO"

            # ----------------------------------------------------
            # FASE 2: Clasificación Temática con IA (Ollama / LLM)
            # ----------------------------------------------------
            if classify_with_ai:
                # Usamos el abstract recién enriquecido, o el resumen original existente
                abstract_para_llm = item.abstract_enriquecido or item.resumen_original
                
                # Para clasificar necesitamos al menos título o abstract
                if item.titulo or abstract_para_llm:
                    articulo_sintetico = MetadataResponse(
                        title=item.titulo or "Sin título",
                        authors=[],
                        doi=item.doi,
                        abstract=abstract_para_llm,
                        keywords=list(keywords_acumuladas) if keywords_acumuladas else None,
                        source="internal"
                    )

                    # 2.1 Clasificación OCDE
                    try:
                        res_ocde = await self.classification_service.classify_metadata_ocde(articulo_sintetico)
                        if res_ocde and res_ocde.clasificaciones:
                            item.area_ocde = [
                                c.model_dump() if hasattr(c, "model_dump") else c.dict()
                                for c in res_ocde.clasificaciones
                            ]
                    except Exception as e_ocde:
                        print(f"Error clasificando OCDE para {item.dspace_uuid}: {e_ocde}")

                    # 2.2 Clasificación ODS
                    try:
                        res_ods = await self.classification_service.classify_metadata_ods(articulo_sintetico)
                        if res_ods and res_ods.clasificaciones:
                            item.area_ods = [
                                c.model_dump() if hasattr(c, "model_dump") else c.dict()
                                for c in res_ods.clasificaciones
                            ]
                    except Exception as e_ods:
                        print(f"Error clasificando ODS para {item.dspace_uuid}: {e_ods}")

                    # Si se obtuvo clasificación con IA
                    if item.area_ocde or item.area_ods:
                        if item.estado == "ENRIQUECIDO_FUENTES":
                            item.estado = "COMPLETO"
                        elif item.estado == "NO_ENCONTRADO":
                            item.estado = "CLASIFICADO_LLM"

            item.mensaje_error = None

        except Exception as e:
            item.estado = "ERROR"
            item.mensaje_error = str(e)
            print(f"Error general enriqueciendo item {item.id} ({item.dspace_uuid}): {e}")

        item.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(item)

        return {
            "id": item.id,
            "dspace_uuid": item.dspace_uuid,
            "doi": item.doi,
            "titulo": item.titulo,
            "estado": item.estado,
            "fuente_enriquecimiento": item.fuente_enriquecimiento,
            "tiene_abstract": bool(item.abstract_enriquecido),
            "abstract_enriquecido": item.abstract_enriquecido,
            "palabras_clave": item.palabras_clave,
            "fuente_financiamiento": item.fuente_financiamiento,
            "derechos_acceso": item.derechos_acceso,
            "area_ocde": item.area_ocde,
            "area_ods": item.area_ods,
            "mensaje_error": item.mensaje_error
        }

    async def enrich_item_by_id(self, item_id: int, db: Session, classify_with_ai: bool = True) -> Optional[Dict[str, Any]]:
        """
        Enriquece un ítem individual buscando por su ID de tabla.
        """
        item = db.query(DSpaceEnrichmentQueue).filter(DSpaceEnrichmentQueue.id == item_id).first()
        if not item:
            return None
        return await self.enrich_item(item, db, classify_with_ai=classify_with_ai)

    async def enrich_item_by_uuid(self, dspace_uuid: str, db: Session, classify_with_ai: bool = True) -> Optional[Dict[str, Any]]:
        """
        Enriquece un ítem individual buscando por su UUID de DSpace.
        """
        item = db.query(DSpaceEnrichmentQueue).filter(DSpaceEnrichmentQueue.dspace_uuid == dspace_uuid).first()
        if not item:
            return None
        return await self.enrich_item(item, db, classify_with_ai=classify_with_ai)

    async def enrich_batch(self, limit: int = 5, estado_filtro: str = "PENDIENTE", classify_with_ai: bool = True, db: Session = None) -> Dict[str, Any]:
        """
        Toma un lote de registros pendientes en la cola y los procesa secuencialmente
        para respetar las políticas de tasa (rate limits) de las fuentes externas.
        """
        items = (
            db.query(DSpaceEnrichmentQueue)
            .filter(DSpaceEnrichmentQueue.estado == estado_filtro)
            .order_by(DSpaceEnrichmentQueue.id.asc())
            .limit(limit)
            .all()
        )

        resultados = []
        for it in items:
            res = await self.enrich_item(it, db, classify_with_ai=classify_with_ai)
            resultados.append(res)
            # Pequeña pausa de cortesía entre llamadas
            await asyncio.sleep(0.5)

        completos = sum(1 for r in resultados if r.get("estado") == "COMPLETO")
        enriquecidos_fuentes = sum(1 for r in resultados if r.get("estado") == "ENRIQUECIDO_FUENTES")
        clasificados_llm = sum(1 for r in resultados if r.get("estado") == "CLASIFICADO_LLM")
        no_encontrados = sum(1 for r in resultados if r.get("estado") == "NO_ENCONTRADO")
        errores = sum(1 for r in resultados if r.get("estado") == "ERROR")

        return {
            "total_procesados": len(resultados),
            "resumen": {
                "completos": completos,
                "enriquecidos_fuentes": enriquecidos_fuentes,
                "clasificados_llm": clasificados_llm,
                "no_encontrados": no_encontrados,
                "errores": errores
            },
            "items": resultados
        }

    async def generate_vriic_sample(
        self,
        db: Session,
        sample_size: int = 35,
        output_filename: str = "muestra_clasificaciones_vriic.xlsx",
        force_reclassify: bool = False
    ) -> Dict[str, Any]:
        """
        Extrae una muestra de publicaciones desde 'dspace_enrichment_queue' que tengan título y resumen original,
        las clasifica en áreas OCDE y ODS usando el LLM (Ollama / LLaMA 3.2),
        persiste las clasificaciones en la base de datos (con su uuid, título y resumen)
        y exporta la muestra formateada en un archivo Excel en la raíz del proyecto para validación con la VRIIC.
        Si force_reclassify=True, ignora clasificaciones previas y vuelve a clasificarlas desde cero con el LLM.
        """
        
        # 1. Extraer publicaciones que tengan título y resumen original válido
        candidatos = (
            db.query(DSpaceEnrichmentQueue)
            .filter(
                DSpaceEnrichmentQueue.titulo.isnot(None),
                func.length(DSpaceEnrichmentQueue.titulo) > 5,
                DSpaceEnrichmentQueue.resumen_original.isnot(None),
                func.length(DSpaceEnrichmentQueue.resumen_original) > 20,
                ~DSpaceEnrichmentQueue.resumen_original.in_(["No disponible", "Sin resumen", "[No abstract available]"])
            )
            # Si forzamos reclasificación, ordenamos por ID; si no, priorizamos las ya clasificadas
            .order_by(
                DSpaceEnrichmentQueue.id.asc() if force_reclassify else
                case(
                    (and_(DSpaceEnrichmentQueue.area_ocde.isnot(None), DSpaceEnrichmentQueue.area_ods.isnot(None)), 0),
                    else_=1
                ),
                DSpaceEnrichmentQueue.id.asc()
            )
            .limit(sample_size)
            .all()
        )

        if not candidatos:
            raise Exception("No se encontraron publicaciones con título y resumen original en la tabla dspace_enrichment_queue.")

        # 2. Clasificar los que aún no tengan clasificación OCDE u ODS (o todos si force_reclassify=True)
        clasificados_nuevos = 0
        ya_clasificados = 0

        for item in candidatos:
            necesita_clasificar_ocde = force_reclassify or not item.area_ocde
            necesita_clasificar_ods = force_reclassify or not item.area_ods

            if necesita_clasificar_ocde or necesita_clasificar_ods:
                articulo_sintetico = MetadataResponse(
                    title=item.titulo or "Sin título",
                    authors=[],
                    doi=item.doi,
                    abstract=item.resumen_original,
                    keywords=item.palabras_clave or None,
                    source="internal"
                )

                if necesita_clasificar_ocde:
                    try:
                        res_ocde = await self.classification_service.classify_metadata_ocde(articulo_sintetico)
                        if res_ocde and res_ocde.clasificaciones:
                            item.area_ocde = [
                                c.model_dump() if hasattr(c, "model_dump") else c.dict()
                                for c in res_ocde.clasificaciones
                            ]
                    except Exception as e_ocde:
                        print(f"Error clasificando OCDE en muestra para {item.dspace_uuid}: {e_ocde}")

                if necesita_clasificar_ods:
                    try:
                        res_ods = await self.classification_service.classify_metadata_ods(articulo_sintetico)
                        if res_ods and res_ods.clasificaciones:
                            item.area_ods = [
                                c.model_dump() if hasattr(c, "model_dump") else c.dict()
                                for c in res_ods.clasificaciones
                            ]
                    except Exception as e_ods:
                        print(f"Error clasificando ODS en muestra para {item.dspace_uuid}: {e_ods}")

                if item.area_ocde and item.area_ods:
                    item.estado = "COMPLETO"
                elif item.area_ocde or item.area_ods:
                    item.estado = "CLASIFICADO_LLM"

                item.updated_at = datetime.utcnow()
                db.commit()
                db.refresh(item)
                clasificados_nuevos += 1
                await asyncio.sleep(0.3)
            else:
                ya_clasificados += 1

        # 3. Generar el archivo Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Muestra VRIIC"
        ws.views.sheetView[0].showGridLines = True

        # Estilos visuales institucionales
        fuente_titulo = Font(name="Calibri", size=14, bold=True, color="002060")
        fuente_subtitulo = Font(name="Calibri", size=10, italic=True, color="595959")
        fuente_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        fuente_datos = Font(name="Calibri", size=9)
        
        fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        fill_header_vriic = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        fill_par = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        fill_impar = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        borde_fino = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        align_center = Alignment(horizontal="center", vertical="top", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Encabezado institucional
        ws.merge_cells("A1:N1")
        ws["A1"] = "MUESTRA DE CLASIFICACIÓN DE PUBLICACIONES (ÁREAS OCDE Y ODS) - VALIDACIÓN VRIIC"
        ws["A1"].font = fuente_titulo
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:N2")
        ws["A2"] = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Modelo: LLaMA 3.2 (Ollama) | Total Muestra: {len(candidatos)} publicaciones | Fuente: Base de datos dspace_enrichment_queue"
        ws["A2"].font = fuente_subtitulo
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

        headers = [
            "N°",
            "UUID DSpace",
            "DOI",
            "Título de la Publicación",
            "Resumen (Abstract Original)",
            "Código Disciplina OCDE",
            "Área / Disciplina OCDE",
            "Justificación OCDE (Modelo)",
            "Código ODS",
            "Nombre ODS",
            "Justificación ODS (Modelo)",
            "Validación VRIIC - OCDE",
            "Validación VRIIC - ODS",
            "Observaciones VRIIC"
        ]

        ws.append([]) # Fila 3 en blanco
        ws.append(headers) # Fila 4
        
        header_row_idx = 4
        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = fuente_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if "VRIIC" in col_name:
                cell.fill = fill_header_vriic
            else:
                cell.fill = fill_header
            cell.border = borde_fino
        
        ws.row_dimensions[header_row_idx].height = 28

        # Llenar datos de cada publicación
        for idx, item in enumerate(candidatos, 1):
            # Parsear OCDE
            cod_ocde_str = ""
            nom_ocde_str = ""
            just_ocde_str = ""
            if item.area_ocde and isinstance(item.area_ocde, list):
                cods, noms, justs = [], [], []
                for c in item.area_ocde:
                    cod = c.get("codigo_disciplina") or c.get("codigo_area") or ""
                    nom = c.get("nombre_disciplina") or c.get("nombre_area") or ""
                    just = c.get("justificacion") or ""
                    if cod: cods.append(str(cod))
                    if nom: noms.append(str(nom))
                    if just: justs.append(str(just))
                cod_ocde_str = "\n".join(cods)
                nom_ocde_str = "\n".join(noms)
                just_ocde_str = "\n---\n".join(justs)

            # Parsear ODS
            num_ods_str = ""
            nom_ods_str = ""
            just_ods_str = ""
            if item.area_ods and isinstance(item.area_ods, list):
                nums, noms, justs = [], [], []
                for o in item.area_ods:
                    num = o.get("numero_ods") or ""
                    nom = o.get("nombre_ods") or ""
                    just = o.get("justificacion") or ""
                    if num: nums.append(f"ODS {num}")
                    if nom: noms.append(str(nom))
                    if just: justs.append(str(just))
                num_ods_str = "\n".join(nums)
                nom_ods_str = "\n".join(noms)
                just_ods_str = "\n---\n".join(justs)

            row_data = [
                idx,
                item.dspace_uuid,
                item.doi or "No registrado",
                item.titulo or "",
                item.resumen_original or "",
                cod_ocde_str,
                nom_ocde_str,
                just_ocde_str,
                num_ods_str,
                nom_ods_str,
                just_ods_str,
                "", # Celda para VRIIC - OCDE
                "", # Celda para VRIIC - ODS
                ""  # Observaciones VRIIC
            ]

            ws.append(row_data)
            current_row = ws.max_row
            fill_actual = fill_par if idx % 2 == 0 else fill_impar

            for c_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=c_idx)
                cell.font = fuente_datos
                cell.fill = fill_actual
                cell.border = borde_fino
                if c_idx in [1, 2, 3, 6, 9]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        # Configurar anchos de columna óptimos
        col_widths = {
            1: 6,    # N°
            2: 38,   # UUID
            3: 24,   # DOI
            4: 40,   # Título
            5: 65,   # Resumen
            6: 18,   # Código OCDE
            7: 30,   # Disciplina OCDE
            8: 45,   # Justificación OCDE
            9: 14,   # N° ODS
            10: 28,  # Nombre ODS
            11: 45,  # Justificación ODS
            12: 24,  # Validación OCDE
            13: 24,  # Validación ODS
            14: 30   # Observaciones
        }

        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Guardar archivo en la raíz del proyecto
        file_path = os.path.abspath(output_filename)
        wb.save(file_path)

        return {
            "status": "success",
            "archivo_generado": output_filename,
            "ruta_absoluta": file_path,
            "total_muestra": len(candidatos),
            "clasificados_nuevos": clasificados_nuevos,
            "ya_clasificados_reutilizados": ya_clasificados,
            "mensaje": f"Muestra de {len(candidatos)} publicaciones generada y guardada exitosamente en {output_filename}"
        }

    def reset_classifications(
        self,
        db: Session,
        item_id: Optional[int] = None,
        dspace_uuid: Optional[str] = None,
        reset_all: bool = False
    ) -> Dict[str, Any]:
        """
        Elimina las clasificaciones (area_ocde y area_ods) de los registros en dspace_enrichment_queue
        para permitir volver a clasificarlos desde cero.
        - Si se especifica item_id o dspace_uuid: resetea ese ítem en particular.
        - Si reset_all=True: resetea todos los registros que tengan áreas OCDE u ODS asignadas.
        """
        query = db.query(DSpaceEnrichmentQueue)
        if item_id:
            query = query.filter(DSpaceEnrichmentQueue.id == item_id)
        elif dspace_uuid:
            query = query.filter(DSpaceEnrichmentQueue.dspace_uuid == dspace_uuid)
        elif reset_all:
            query = query.filter(
                (DSpaceEnrichmentQueue.area_ocde.isnot(None)) | 
                (DSpaceEnrichmentQueue.area_ods.isnot(None))
            )
        else:
            return {
                "status": "error",
                "mensaje": "Debe especificar 'item_id', 'dspace_uuid' o indicar 'reset_all=True' para reiniciar las clasificaciones."
            }

        items = query.all()
        if not items:
            return {
                "status": "info",
                "total_reseteados": 0,
                "mensaje": "No se encontraron registros con clasificaciones para reiniciar."
            }

        for item in items:
            item.area_ocde = None
            item.area_ods = None
            if item.abstract_enriquecido:
                item.estado = "ENRIQUECIDO_FUENTES"
            else:
                item.estado = "PENDIENTE"
            item.updated_at = datetime.utcnow()

        db.commit()

        return {
            "status": "success",
            "total_reseteados": len(items),
            "mensaje": f"Se eliminaron las clasificaciones de {len(items)} publicación(es) y se restableció su estado."
        }
